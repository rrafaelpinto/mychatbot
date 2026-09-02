from pathlib import Path

import pymupdf
from django.core.files.uploadedfile import UploadedFile
from django.utils.text import get_valid_filename
from docx import Document

from chatbot.models import CandidateDocument, CandidateProfile
from chatbot.validators import DOCX_CONTENT_TYPE, PDF_CONTENT_TYPE, detect_document_type

MAX_PDF_PAGES = 50
MAX_EXTRACTED_CHARACTERS = 200_000


class DocumentProcessingError(Exception):
    """A source document could not be safely processed."""


def create_candidate_document(
    candidate: CandidateProfile, uploaded_file: UploadedFile
) -> CandidateDocument:
    detected_type = detect_document_type(uploaded_file)
    original_name = get_valid_filename(Path(uploaded_file.name).name)[:255]
    document = CandidateDocument(
        candidate=candidate,
        file=uploaded_file,
        original_filename=original_name,
        detected_content_type=detected_type,
        size_bytes=uploaded_file.size,
    )
    document.full_clean()
    document.save()
    return document


def process_candidate_document(document: CandidateDocument) -> None:
    try:
        with document.file.open("rb") as source:
            if document.detected_content_type == PDF_CONTENT_TYPE:
                text = _extract_pdf(source.read())
            elif document.detected_content_type == DOCX_CONTENT_TYPE:
                text = _extract_docx(source)
            else:
                raise DocumentProcessingError("Unsupported document type")
        document.mark_processed(text[:MAX_EXTRACTED_CHARACTERS])
        document.save(
            update_fields=["extracted_text", "processing_status", "processed_at"]
        )
    except Exception as exc:
        document.mark_failed()
        document.save(
            update_fields=["extracted_text", "processing_status", "processed_at"]
        )
        raise DocumentProcessingError("The document could not be processed.") from exc


def _extract_pdf(content: bytes) -> str:
    with pymupdf.open(stream=content, filetype="pdf") as pdf:
        if len(pdf) > MAX_PDF_PAGES:
            raise DocumentProcessingError("PDF page limit exceeded")
        return "\n".join(page.get_text() for page in pdf)


def _extract_docx(source) -> str:
    document = Document(source)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)
